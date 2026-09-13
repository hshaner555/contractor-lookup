"""
Ohio OCILB collector/importer.

Two supported paths:

1) Bulk roster import (preferred when Ohio's generated roster is available).
2) Public detail-page refresh using PrintLicenseDetails.aspx URLs.

The Ohio site is ASP.NET/session based, so this module deliberately avoids
hard-coded __VIEWSTATE / __EVENTVALIDATION values.

Official public detail pages expose:
- contractor/person name
- business/company
- credential number
- license type
- issue date
- expiration date
- status/reason
"""

from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from contractor_importers.persistence import upsert_record
from contractor_lib.models import CredentialRecord
from openpyxl import load_workbook

AUTHORITY = "Ohio Construction Industry Licensing Board (OCILB)"
JURISDICTION = "State of Ohio"

DETAIL_URL = "https://elicense4.com.ohio.gov/Lookup/PrintLicenseDetails.aspx"

HEADER_ALIASES = {
    "credential": "credential_number",
    "credential number": "credential_number",
    "license number": "credential_number",
    "license type": "credential_type",
    "type": "credential_type",
    "name": "person_name",
    "licensee name": "person_name",
    "company": "business_name",
    "company name": "business_name",
    "business name": "business_name",
    "expiration date": "expiration_date",
    "expiration": "expiration_date",
    "status": "status",
    "issue date": "issue_date",
    "county": "county",
    "state": "state",
}


def _clean(v):
    if v is None:
        return None
    v = re.sub(r"\s+", " ", str(v)).strip()
    return v or None


def _norm_header(v):
    return re.sub(r"\s+", " ", (v or "").replace("\n", " ").strip().lower())


def _row_to_record(row: dict, source_url: str | None = None):
    number = _clean(row.get("credential_number"))
    ctype = _clean(row.get("credential_type")) or "Contractor"
    return CredentialRecord(
        business_name=_clean(row.get("business_name")),
        person_name=_clean(row.get("person_name")),
        credential_number=number,
        credential_type=ctype,
        credential_kind="LICENSE",
        issuing_authority=AUTHORITY,
        jurisdiction=JURISDICTION,
        status=_clean(row.get("status")),
        expiration_date=_clean(row.get("expiration_date")),
        source_url=source_url,
        last_verified=datetime.now(timezone.utc).date().isoformat(),
    )


def parse_roster_rows(rows):
    out = []
    for raw in rows:
        mapped = {}
        for k, v in raw.items():
            key = HEADER_ALIASES.get(_norm_header(k))
            if key:
                mapped[key] = _clean(v)
        if mapped.get("credential_number"):
            out.append(_row_to_record(mapped))
    return out


def parse_roster_file(path: str):
    """Parse an Ohio-generated .xlsx, .csv, or tab-delimited .txt roster."""
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix == ".xlsx":
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            raise ValueError("Ohio roster workbook has no active worksheet")
        values = ws.iter_rows(values_only=True)
        headers = [str(x or "") for x in next(values)]
        rows = [dict(zip(headers, row)) for row in values]
        return parse_roster_rows(rows)

    if suffix in {".csv", ".txt"}:
        read_text = p.read_text(encoding="utf-8-sig", errors="replace")
        sample = read_text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t|")
        except csv.Error:
            dialect = csv.excel_tab if "\t" in sample else csv.excel
        reader = csv.DictReader(read_text.splitlines(), dialect=dialect)
        return parse_roster_rows(reader)

    raise ValueError("Supported Ohio roster formats: .xlsx, .csv, .txt")


def ingest_roster_file(path: str):
    result = {"records": 0, "inserted": 0, "updated": 0}
    for rec in parse_roster_file(path):
        result["records"] += 1
        action = str(upsert_record(rec)["action"])
        result[action] += 1
    return result


def parse_detail_html(html: str, source_url: str | None = None):
    """
    Parse an official OCILB PrintLicenseDetails page.

    The page can contain multiple credentials for one person/company. We return
    one CredentialRecord per credential so HVAC + Hydronics, etc. remain
    separate credentials under the same contractor profile.
    """
    soup = BeautifulSoup(html, "lxml")
    _ = soup.get_text(" ", strip=True)

    person_name = None
    # Prefer the first table headed by Name / Company / Public Address.
    for table in soup.find_all("table"):
        rows = table.find_all("tr")
        if not rows:
            continue
        headers = [_norm_header(x.get_text(" ", strip=True)) for x in rows[0].find_all(["th", "td"])]
        if "name" in headers and ("company" in headers or "public address" in headers):
            if len(rows) > 1:
                vals = [_clean(x.get_text(" ", strip=True)) for x in rows[1].find_all(["th", "td"])]
                row = dict(zip(headers, vals))
                person_name = row.get("name")
            break

    records = []
    for table in soup.find_all("table"):
        rows = table.find_all("tr")
        if not rows:
            continue
        headers = [_norm_header(x.get_text(" ", strip=True)) for x in rows[0].find_all(["th", "td"])]
        if "credential" not in headers or "license type" not in headers:
            continue

        for tr in rows[1:]:
            vals = [_clean(x.get_text(" ", strip=True)) for x in tr.find_all(["th", "td"])]
            if not vals:
                continue
            row = dict(zip(headers, vals))
            number = row.get("credential")
            if not number:
                continue
            records.append(
                CredentialRecord(
                    business_name=row.get("company"),
                    person_name=person_name,
                    credential_number=number,
                    credential_type=row.get("license type") or "Contractor",
                    credential_kind="LICENSE",
                    issuing_authority=AUTHORITY,
                    jurisdiction=JURISDICTION,
                    status=row.get("status"),
                    expiration_date=row.get("expiration date"),
                    source_url=source_url,
                    last_verified=datetime.now(timezone.utc).date().isoformat(),
                )
            )

    return records


def fetch_detail(contact: str | int, cred: str | int, session: requests.Session | None = None):
    session = session or requests.Session()
    params = {"contact": str(contact), "cred": str(cred)}
    r = session.get(DETAIL_URL, params=params, timeout=30)
    r.raise_for_status()
    return parse_detail_html(r.text, source_url=r.url)


def ingest_detail(contact: str | int, cred: str | int, session: requests.Session | None = None):
    out = {"records": 0, "inserted": 0, "updated": 0}
    for rec in fetch_detail(contact, cred, session=session):
        out["records"] += 1
        action = str(upsert_record(rec)["action"])
        out[action] += 1
    return out
